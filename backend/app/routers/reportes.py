from datetime import date, datetime, time

from fastapi import APIRouter, Depends
from sqlalchemy import func
from sqlalchemy.orm import Session

from ..database import obtener_db
from ..models.venta import Venta
from ..models.usuario import Usuario
from ..security import requerir_admin

from io import BytesIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter



router = APIRouter(
    prefix="/reportes",
    tags=["Reportes"]
)


# =========================================================
# RESUMEN DE VENTAS
# Solo administrador
# =========================================================
@router.get("/ventas/resumen")
def obtener_resumen_ventas(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    usuario_id: int | None = None,
    db: Session = Depends(obtener_db),
    usuario_actual: Usuario = Depends(requerir_admin)
):
    query = db.query(Venta)

    # Filtrar por vendedor
    if usuario_id is not None:
        query = query.filter(
            Venta.usuario_id == usuario_id
        )

    # Filtrar desde una fecha
    if fecha_desde is not None:
        inicio = datetime.combine(
            fecha_desde,
            time.min
        )

        query = query.filter(
            Venta.fecha >= inicio
        )

    # Filtrar hasta una fecha
    if fecha_hasta is not None:
        fin = datetime.combine(
            fecha_hasta,
            time.max
        )

        query = query.filter(
            Venta.fecha <= fin
        )

    # Cantidad de ventas
    cantidad_ventas = query.count()

    # Total vendido
    total_vendido = (
        query.with_entities(
            func.coalesce(
                func.sum(Venta.total),
                0
            )
        ).scalar()
    )

    return {
        "cantidad_ventas": cantidad_ventas,
        "total_vendido": total_vendido
    }
    
# =========================================================
# EXPORTAR VENTAS A EXCEL
# Solo administrador
# =========================================================
@router.get("/ventas/excel")
def exportar_ventas_excel(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    usuario_id: int | None = None,
    db: Session = Depends(obtener_db),
    usuario_actual: Usuario = Depends(requerir_admin)
):
    query = db.query(Venta)

    # Filtrar por vendedor
    if usuario_id is not None:
        query = query.filter(
            Venta.usuario_id == usuario_id
        )

    # Filtrar desde
    if fecha_desde is not None:
        inicio = datetime.combine(
            fecha_desde,
            time.min
        )

        query = query.filter(
            Venta.fecha >= inicio
        )

    # Filtrar hasta
    if fecha_hasta is not None:
        fin = datetime.combine(
            fecha_hasta,
            time.max
        )

        query = query.filter(
            Venta.fecha <= fin
        )

    ventas = (
        query
        .order_by(Venta.fecha.desc())
        .all()
    )

    # Crear Excel
    workbook = Workbook()

    hoja = workbook.active
    hoja.title = "Ventas"

    # Encabezados
    encabezados = [
        "Venta",
        "Fecha",
        "Cliente",
        "Vendedor",
        "Producto",
        "Cantidad",
        "Precio unitario",
        "Subtotal",
        "Total venta",
    ]

    hoja.append(encabezados)

    # Estilo encabezados
    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )

        celda.alignment = Alignment(
            horizontal="center"
        )

    # Datos
    for venta in ventas:

        cliente_nombre = (
            venta.cliente.nombre
            if venta.cliente
            else "Consumidor final"
        )

        vendedor_nombre = (
            venta.usuario.nombre
            if venta.usuario
            else "Sin vendedor"
        )

        for detalle in venta.detalles:

            hoja.append([
                venta.id,
                venta.fecha,
                cliente_nombre,
                vendedor_nombre,
                detalle.producto.nombre,
                detalle.cantidad,
                float(detalle.precio_unitario),
                float(detalle.subtotal),
                float(venta.total),
            ])

    # Formato monetario
    for fila in range(2, hoja.max_row + 1):
        hoja.cell(
            row=fila,
            column=7
        ).number_format = '$ #,##0.00'

        hoja.cell(
            row=fila,
            column=8
        ).number_format = '$ #,##0.00'

        hoja.cell(
            row=fila,
            column=9
        ).number_format = '$ #,##0.00'

    # Formato fecha
    for fila in range(2, hoja.max_row + 1):
        hoja.cell(
            row=fila,
            column=2
        ).number_format = "dd/mm/yyyy hh:mm"

    # Ajustar ancho de columnas
    for columna in hoja.columns:

        largo_maximo = 0

        letra = get_column_letter(
            columna[0].column
        )

        for celda in columna:

            if celda.value is not None:
                largo_maximo = max(
                    largo_maximo,
                    len(str(celda.value))
                )

        hoja.column_dimensions[letra].width = min(
            largo_maximo + 3,
            35
        )

    # Congelar encabezado
    hoja.freeze_panes = "A2"

    # Filtro de Excel
    hoja.auto_filter.ref = hoja.dimensions

    # Guardar en memoria
    archivo = BytesIO()

    workbook.save(archivo)

    archivo.seek(0)

    nombre_archivo = (
        f"reporte_ventas_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    return StreamingResponse(
        archivo,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{nombre_archivo}"'
            )
        }
    )